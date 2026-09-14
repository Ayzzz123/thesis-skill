# -*- coding: utf-8 -*-
"""Test B：构建 + QA（正例全过 / 负例逐项 FAIL / 校模提取 / md 输入源 /
标点全半角 / 附录页隐藏放映 / 逐字讲稿导出）

用法: python ppt-direct/tests/test_b_build_and_qa.py
退出码: 0=全过；1=有失败
"""
import json
import os
import shutil
import sys
import tempfile

import yaml

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
sys.path.insert(0, os.path.join(ROOT, "scripts"))
import build_pptx
import ingest_source
import ppt_qa
import speech_export
import theme_extract
from pptx import Presentation

THEME = os.path.join(ROOT, "assets", "theme-default.yaml")
PASS, FAIL = 0, 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS {name}")
    else:
        FAIL += 1
        print(f"  FAIL {name} | {detail}")


def write_deck(root, slides, tier="short"):
    deck = {"meta": {"title": "测试题目", "presenter": "张三",
                     "major": "飞行器维修工程技术", "page_tier": tier},
            "slides": slides}
    path = os.path.join(root, "deck.yaml")
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(deck, f, allow_unicode=True)
    return path


def good_slides():
    return [
        {"layout": "cover", "notes":
            "各位老师好，我是来自飞行器维修工程技术专业的张三。今天我汇报的题目是"
            "《测试题目》，汇报将从研究背景、研究方法、分析过程与结果、结论与建议"
            "四个方面依次展开，请各位老师批评指正。"},
        {"layout": "toc", "notes":
            "本次汇报共分为四个部分。第一部分介绍研究背景与工程实际意义，第二部分"
            "说明研究方法与技术路线，第三部分展示分析过程与主要结果，第四部分给出"
            "结论与改进建议。下面首先进入第一部分。"},
        {"layout": "section", "no": 1, "title": "研究背景", "notes":
            "首先进入第一部分研究背景。该部分主要说明本课题的工程实际背景、当前"
            "存在的问题与局限性，以及开展本研究的现实意义，为后续章节奠定基础。"},
        {"layout": "content", "title": "背景",
         "bullets": ["要点一", "要点二", [1, "二级要点"]], "notes":
            "本页介绍研究背景的核心内容。当前该问题在工程实际中较为突出，现有"
            "检测与维护手段存在一定局限，因此有必要开展针对性的分析与研究，从而"
            "为后续的维修决策和改进措施提供可靠依据。"},
        {"layout": "two_column", "title": "方法对比",
         "left": {"title": "FMEA", "bullets": ["正向枚举"]},
         "right": {"title": "FTA", "bullets": ["逆向追溯"]},
         "notes":
            "在方法选择上，本页对比了两种常用分析手段。FMEA 从故障模式出发正向"
            "枚举影响与危害度，FTA 从顶事件出发逆向追溯失效原因链，二者各有适用"
            "场景，本课题将两者结合使用以相互补充验证。"},
        {"layout": "section", "no": 2, "title": "结果", "notes":
            "接下来进入第二部分结果分析。本部分将展示主要的分析过程、关键数据"
            "以及对比验证结果，并对结果所反映的规律和工程含义进行简要讨论。"},
        {"layout": "content", "title": "结果",
         "bullets": ["发现一", "发现二"], "notes":
            "本页呈现了主要分析结果。从结果可以看出，关键结论与预期基本一致，"
            "相关数据经过交叉验证具有一致性，可靠性水平满足工程要求，能够为后续"
            "的维修策略改进提供直接支撑。"},
        {"layout": "closing", "notes":
            "以上就是我汇报的全部内容，感谢各位老师的聆听。本课题还存在一些不足"
            "之处，后续我会继续完善。恳请各位老师对论文和本课题提出宝贵的指导"
            "意见，谢谢大家。"},
    ]


def build_and_qa(root, deck_path):
    pptx = os.path.join(root, "out.pptx")
    layout = os.path.join(root, "layout.json")
    build_pptx.build(deck_path, THEME, pptx, layout)
    qa_dir = os.path.join(root, "qa")
    code = ppt_qa.run_qa(pptx, layout, THEME, deck_path, qa_dir)
    report = open(os.path.join(qa_dir, "ppt-qa-report.md"),
                  encoding="utf-8").read()
    return code, report, pptx, layout


def main():
    tmp = tempfile.mkdtemp(prefix="ppd_build_")
    try:
        # --- 正例：全部 PASS ---
        d1 = os.path.join(tmp, "good")
        os.makedirs(d1)
        deck = write_deck(d1, good_slides())
        code, report, pptx, layout = build_and_qa(d1, deck)
        check("正例 QA 全过", code == 0, report)
        check("pptx 生成", os.path.getsize(pptx) > 5000)
        with open(layout, encoding="utf-8") as f:
            lj = json.load(f)
        check("layout sidecar 8 页", len(lj["pages"]) == 8)
        toc_body = next(s for s in lj["pages"][1]["shapes"]
                        if s.get("kind") == "body")
        check("toc 自动收集板块", toc_body["paras"] == 2,
              f"toc 段落数 {toc_body['paras']} != 2 个板块")

        # --- 负例：要点超限 + 占位符 + 缺备注 ---
        d2 = os.path.join(tmp, "bad")
        os.makedirs(d2)
        bad = good_slides()
        bad[3] = {"layout": "content", "title": "【待填】",
                  "bullets": [f"第{i}条要点内容" for i in range(9)],
                  "notes": ""}
        deck2 = write_deck(d2, bad)
        code, report, _, _ = build_and_qa(d2, deck2)
        check("负例 QA 退出码 1", code == 1)
        check("PPT-02 命中", "PPT-02: FAIL" in report)
        check("PPT-08 命中占位符", "PPT-08: FAIL" in report)
        check("PPT-10 命中缺备注", "PPT-10: FAIL" in report)

        # --- 负例：超长文本触发溢出估算（自动缩字到 14pt 触底仍溢出） ---
        d3 = os.path.join(tmp, "overflow")
        os.makedirs(d3)
        ov = good_slides()
        ov[3] = {"layout": "content", "title": "溢出页",
                 "bullets": ["长" * 3000], "notes": "n"}
        code, report, _, _ = build_and_qa(d3, write_deck(d3, ov))
        check("PPT-04 命中溢出", "PPT-04: FAIL" in report, report)
        check("PPT-03 命中字数", "PPT-03: FAIL" in report)
        check("PPT-04 附缩字号自愈说明", "自动缩字号" in report)

        # --- 页数档位 ---
        d4 = os.path.join(tmp, "tier")
        os.makedirs(d4)
        code, report, _, _ = build_and_qa(
            d4, write_deck(d4, good_slides()[:5], tier="long"))
        check("PPT-01 命中页数不足", "PPT-01: FAIL" in report)

        # --- 校模提取（用正例 pptx 当模板） ---
        theme_out = os.path.join(tmp, "theme-school.yaml")
        info = theme_extract.extract(pptx)
        check("提取画幅 16:9", info["size"] == "16:9")
        check("提取到中文字体", bool(info["fonts"]["cjk"]))
        check("提取到主题色", "primary" in info["colors"])

        # --- md 输入源 ---
        md = os.path.join(tmp, "论文.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write("# 绪论\n\n研究背景段落内容，超过十二个字的句子。\n\n"
                    "# 结论\n\n结论段落内容，也是超过十二个字的句子。\n")
        out_dir = os.path.join(tmp, "pptproj")
        argv = sys.argv
        sys.argv = ["ingest_source.py", "--md", md, "--out", out_dir]
        try:
            rc = ingest_source.main()
        finally:
            sys.argv = argv
        check("ingest 退出码 0", rc == 0)
        deck_draft = os.path.join(out_dir, ".pptdirect", "artifacts",
                                  "deck.yaml")
        with open(deck_draft, encoding="utf-8") as f:
            draft = yaml.safe_load(f)
        check("ingest 产出 deck 草稿",
              any(s["layout"] == "section" for s in draft["slides"]))
        check("ingest 保留占位符纪律",
              draft["meta"]["presenter"] == "【待填】")

        # --- 表格版式：正例 QA 全过 + PPT-12 超规负例 ---
        d5 = os.path.join(tmp, "table")
        os.makedirs(d5)
        tbl = good_slides()
        tbl[3] = {"layout": "table", "title": "FMEA 关键行",
                  "table": {"headers": ["故障模式", "RPN"],
                            "rows": [[f"模式{i}", str(i * 10)] for i in range(4)]},
                  "notes":
                    "本页以表格呈现 FMEA 排序前四的故障模式与 RPN 值。可以看出风险"
                    "主要集中在少数几项模式上，这正是后续维修策略需要优先关注的"
                    "对象，具体改进措施将在下一板块展开说明。"}
        code, report, _, _ = build_and_qa(d5, write_deck(d5, tbl))
        check("表格正例 QA 全过", code == 0, report)
        d5b = os.path.join(tmp, "table_bad")
        os.makedirs(d5b)
        tbl_bad = good_slides()
        tbl_bad[3] = {"layout": "table", "title": "超规表",
                      "table": {"headers": ["列1", "列2"],
                                "rows": [[f"行{i}", "x"] for i in range(13)]},
                      "notes": "表格页"}
        code, report, _, _ = build_and_qa(d5b, write_deck(d5b, tbl_bad))
        check("PPT-12 命中表格超规", "PPT-12: FAIL" in report, report)

        # --- 窄栏自动缩字号（two_column 长文本） ---
        d6 = os.path.join(tmp, "shrink")
        os.makedirs(d6)
        sh = good_slides()
        sh[4] = {"layout": "two_column", "title": "窄栏压字",
                 "left": {"title": "FMEA", "bullets": ["长文本" * 150]},
                 "right": {"title": "FTA", "bullets": ["逆向追溯"]},
                 "notes":
                    "窄栏里的超长文本会触发引擎自动缩字号，直到内容放得下为止。"
                    "这是度量级拟合的自愈能力验证页，缩字号事件会记录在 sidecar "
                    "中供 QA 复核。"}
        code, report, _, layout6 = build_and_qa(d6, write_deck(d6, sh))
        with open(layout6, encoding="utf-8") as f:
            lj6 = json.load(f)
        shrunk = [(p["page"], s["shrink_from"], s["font_pt"])
                  for p in lj6["pages"] for s in p["shapes"]
                  if s.get("shrink_from")]
        check("窄栏触发自动缩字号", bool(shrunk), str(shrunk))

        # --- PPT-11 负例：deck 声明 duration_min 远超实际讲稿 ---
        d7 = os.path.join(tmp, "dur")
        os.makedirs(d7)
        deck7 = write_deck(d7, good_slides())
        with open(deck7, encoding="utf-8") as f:
            dd7 = yaml.safe_load(f)
        dd7["meta"]["duration_min"] = 20
        with open(deck7, "w", encoding="utf-8") as f:
            yaml.safe_dump(dd7, f, allow_unicode=True)
        code, report, _, _ = build_and_qa(d7, deck7)
        check("PPT-11 命中时长不足", "PPT-11: FAIL" in report, report)

        # --- 模板模式：中文版式名 + PPT-13 使用率 ---
        d8 = os.path.join(tmp, "tpl")
        os.makedirs(d8)
        tpl_prs = Presentation()
        for i, name in {0: "标题幻灯片", 1: "标题和内容",
                        2: "节标题", 6: "空白"}.items():
            tpl_prs.slide_layouts[i].name = name
        tpl_path = os.path.join(d8, "校模.pptx")
        tpl_prs.save(tpl_path)
        theme_t = yaml.safe_load(open(THEME, encoding="utf-8"))
        theme_t["template_layouts"] = theme_extract.scan_template_layouts(tpl_path)
        theme_path = os.path.join(d8, "theme-t.yaml")
        with open(theme_path, "w", encoding="utf-8") as f:
            yaml.safe_dump(theme_t, f, allow_unicode=True, sort_keys=False)
        deck8 = write_deck(d8, good_slides())
        pptx8 = os.path.join(d8, "out.pptx")
        layout8 = os.path.join(d8, "layout.json")
        build_pptx.build(deck8, theme_path, pptx8, layout8, template=tpl_path)
        with open(layout8, encoding="utf-8") as f:
            lj8 = json.load(f)
        tpl_modes = [p["mode"] for p in lj8["pages"]
                     if p["mode"].startswith("template:")]
        check("模板模式母版驱动 ≥4 页", len(tpl_modes) >= 4, str(tpl_modes))
        code = ppt_qa.run_qa(pptx8, layout8, theme_path, deck8,
                             os.path.join(d8, "qa"))
        report8 = open(os.path.join(d8, "qa", "ppt-qa-report.md"),
                       encoding="utf-8").read()
        check("PPT-13 母版使用率 PASS", "PPT-13: PASS" in report8, report8)

        # --- S10 联动：defense 产物 → 板块重排 + 问答备份页 ---
        d9 = os.path.join(tmp, "s10")
        root9 = os.path.join(d9, "thesis")
        ch9 = os.path.join(root9, ".aeromech", "artifacts", "chapters")
        os.makedirs(ch9)
        open(os.path.join(ch9, "ch1.md"), "w", encoding="utf-8").write(
            "# 第一章 绪论\n\n研究背景与意义段落，这里是超过十二个字的句子内容。\n")
        open(os.path.join(ch9, "ch2.md"), "w", encoding="utf-8").write(
            "# 第二章 基于 FMEA 的故障识别与分析方法\n\n## 方法概述\n\n"
            "分析过程第一步，这是超过十二个字的句子。\n")
        open(os.path.join(ch9, "ch3.md"), "w", encoding="utf-8").write(
            "# 第三章 维修策略优化与建议\n\n## 优化方向\n\n"
            "策略建议段落，这是超过十二个字的句子。\n")
        def9 = os.path.join(root9, ".aeromech", "artifacts", "defense")
        os.makedirs(def9)
        open(os.path.join(def9, "ppt-structure.md"), "w", encoding="utf-8").write(
            "1. 方案与建议\n2. 研究背景与意义\n3. 分析过程与结果\n")
        open(os.path.join(def9, "qa-bank.md"), "w", encoding="utf-8").write(
            "| 类别 | 示例问题 | 推荐回答逻辑 | 回答依据 | 可能追问 |\n"
            "|---|---|---|---|---|\n"
            "| 方法类 | 为什么选择 FMEA？ | 适合枚举功能失效模式 | 第二章 | 局限如何克服？ |\n")
        out9 = os.path.join(d9, "pptproj")
        argv = sys.argv
        sys.argv = ["ingest_source.py", "--aeromech", root9, "--out", out9,
                    "--tier", "standard"]
        try:
            rc = ingest_source.main()
        finally:
            sys.argv = argv
        check("S10 ingest 退出码 0", rc == 0)
        deck9 = yaml.safe_load(open(
            os.path.join(out9, ".pptdirect", "artifacts", "deck.yaml"),
            encoding="utf-8"))
        secs9 = [s["title"] for s in deck9["slides"]
                 if s["layout"] == "section"]
        check("S10 板块重排生效",
              secs9[0] == "方案与建议" and secs9[1] == "研究背景与意义",
              str(secs9))
        check("S10 问答备份页生成",
              any(s.get("appendix") for s in deck9["slides"]))
        check("S10 附录页数写入 meta",
              deck9["meta"].get("appendix_pages", 0) >= 1)

        # --- PPT-14 标点/全半角负例 ---
        d10 = os.path.join(tmp, "p14")
        os.makedirs(d10)
        p14 = good_slides()
        p14[3] = {"layout": "content", "title": "标点页",
                  "bullets": ["第一点,第二点"], "notes": "标点负例"}
        code, report, _, _ = build_and_qa(d10, write_deck(d10, p14))
        check("PPT-14 命中半角标点", "PPT-14: FAIL" in report, report)

        # --- PPT-15 附录页隐藏放映 + speech_export（用 S10 deck 真附录） ---
        d11 = os.path.join(tmp, "s10build")
        os.makedirs(d11)
        deck9_path = os.path.join(out9, ".pptdirect", "artifacts", "deck.yaml")
        code, report, _, layout11 = build_and_qa(d11, deck9_path)
        check("PPT-15 隐藏放映 PASS", "PPT-15: PASS" in report, report)
        with open(layout11, encoding="utf-8") as f:
            lj11 = json.load(f)
        app_pages = [p for p in lj11["pages"] if p.get("appendix")]
        check("sidecar 附录页 hidden 标记",
              bool(app_pages) and all(p.get("hidden") for p in app_pages),
              str([(p["page"], p.get("hidden")) for p in app_pages]))

        speech_path = os.path.join(d11, "speech.md")
        speech_export.export(deck9_path, speech_path)
        speech_md = open(speech_path, encoding="utf-8").read()
        check("speech.md 含时长估算",
              "（约" in speech_md and "分钟）" in speech_md, speech_md[:200])
        check("speech.md 标注附录放映隐藏",
              "【备份页·放映隐藏】" in speech_md, speech_md[:200])

        # --- flow 技术路线图版式：正例 QA 全过 + sidecar 步数 ---
        d12 = os.path.join(tmp, "flow")
        os.makedirs(d12)
        fl = good_slides()
        fl[3] = {"layout": "flow", "title": "技术路线",
                 "steps": ["功能结构分析", "故障树建模", "FMEA 排序", "维修策略"],
                 "notes":
                    "技术路线分四步：先做功能结构分析厘清系统组成，再建立故障树"
                    "追溯失效原因链，接着用 FMEA 对故障模式排序，最后据此提出"
                    "维修策略优化建议，四步环环相扣形成完整分析闭环。"}
        code, report, _, layout12 = build_and_qa(d12, write_deck(d12, fl))
        check("flow 正例 QA 全过", code == 0, report)
        with open(layout12, encoding="utf-8") as f:
            lj12 = json.load(f)
        flow_body = [s for p in lj12["pages"] for s in p["shapes"]
                     if s.get("kind") == "body" and s.get("paras")]
        check("flow sidecar 记录步数", bool(flow_body), str(flow_body))

        # --- PPT-16 数据溯源：正例数字在源 / 负例查无此数 / 无源跳过 ---
        d13 = os.path.join(tmp, "p16")
        os.makedirs(d13)
        src_file = os.path.join(d13, "源.md")
        with open(src_file, "w", encoding="utf-8") as f:
            f.write("# 第一章\n\n实测 RPN 值为 216，可靠度 0.96。\n")
        good_p16 = good_slides()
        good_p16[3] = {"layout": "content", "title": "结果",
                       "bullets": ["RPN 值为 216", "可靠度 0.96"],
                       "notes": "数据溯源正例"}
        deck13 = write_deck(d13, good_p16)
        code, report, pptx13, layout13 = build_and_qa(d13, deck13)
        check("PPT-16 无源时跳过", "PPT-16: PASS" in report, report)
        code = ppt_qa.run_qa(pptx13, layout13, THEME, deck13,
                             os.path.join(d13, "qa"), source_path=src_file)
        report16 = open(os.path.join(d13, "qa", "ppt-qa-report.md"),
                        encoding="utf-8").read()
        check("PPT-16 正例数字在源 PASS", "PPT-16: PASS" in report16, report16)
        bad_p16 = good_slides()
        bad_p16[3] = {"layout": "content", "title": "结果",
                      "bullets": ["RPN 值为 99999"], "notes": "数据溯源负例"}
        deck13b = write_deck(d13, bad_p16)
        code, report, pptx13b, layout13b = build_and_qa(d13, deck13b)
        code = ppt_qa.run_qa(pptx13b, layout13b, THEME, deck13b,
                             os.path.join(d13, "qa"), source_path=src_file)
        report16b = open(os.path.join(d13, "qa", "ppt-qa-report.md"),
                         encoding="utf-8").read()
        check("PPT-16 负例查无此数 FAIL",
              "PPT-16: FAIL" in report16b and "99999" in report16b, report16b)

        # --- docx 输入源自动抽图 → materials/figNN ---
        d14 = os.path.join(tmp, "docx")
        os.makedirs(d14)
        from PIL import Image as PILImage
        from docx import Document
        from docx.shared import Inches
        img1 = os.path.join(d14, "f1.png")
        img2 = os.path.join(d14, "f2.png")
        PILImage.new("RGB", (320, 240), (30, 144, 255)).save(img1)
        PILImage.new("RGB", (320, 240), (255, 80, 80)).save(img2)
        doc = Document()
        doc.add_heading("第一章 绪论", level=1)
        doc.add_paragraph("研究背景段落内容，这是超过十二个字的句子。")
        doc.add_picture(img1, width=Inches(2))
        doc.add_picture(img2, width=Inches(2))
        docx_path = os.path.join(d14, "论文.docx")
        doc.save(docx_path)
        out14 = os.path.join(d14, "pptproj")
        argv = sys.argv
        sys.argv = ["ingest_source.py", "--docx", docx_path, "--out", out14]
        try:
            rc = ingest_source.main()
        finally:
            sys.argv = argv
        mat14 = os.path.join(out14, "materials")
        check("docx 抽图退出码 0", rc == 0)
        check("docx 抽图 materials/fig01+fig02",
              os.path.isfile(os.path.join(mat14, "fig01.png")) and
              os.path.isfile(os.path.join(mat14, "fig02.png")),
              str(sorted(os.listdir(mat14)) if os.path.isdir(mat14) else []))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print(f"\n结果: {PASS} PASS / {FAIL} FAIL")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
