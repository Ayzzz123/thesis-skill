# -*- coding: utf-8 -*-
"""pptx_engine.py — ppt-direct 渲染引擎（python-pptx 直出 .pptx）

版式原语：cover / toc / section / content / two_column / image_text / closing。
所有页面统一从空白版式绘制，主题（配色/字体/字号）全部来自 theme.yaml，
禁止在引擎内硬编码颜色与字体。

绘制即输出 layout JSON sidecar（每页每个文本框的几何与文字量元数据，
文字溢出 est 由引擎按字符宽度估算），供 scripts/ppt_qa.py 消费——
与 aeromech-thesis 的 figkit → graph_quality_qa 同一思路。

用法:
    from pptx_engine import DeckBuilder, load_theme
    theme = load_theme("assets/theme-default.yaml")
    d = DeckBuilder(theme)
    d.add_cover(title=..., presenter=..., ...)
    d.add_content(title=..., bullets=[...], notes=...)
    d.save("答辩PPT.pptx", layout_json="layout.json")
"""
import json
import math
import os

import yaml
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

CM_PER_PT = 1 / 28.3465  # 1pt = 0.3528mm

SIZES = {"16:9": (33.867, 19.05), "4:3": (25.4, 19.05)}

# 占位符残留标记（PPT-08 与内容模块共用此约定）
PLACEHOLDER_MARK = "【待填】"


def load_theme(path):
    with open(path, "r", encoding="utf-8") as f:
        theme = yaml.safe_load(f)
    for k in ("colors", "fonts", "sizes"):
        if k not in theme:
            raise ValueError(f"theme 缺少必填段: {k} ({path})")
    return theme


def _rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _set_run_font(run, theme, size_pt, bold=False, color="text"):
    """同时设置拉丁与东亚字体（python-pptx 默认只写 latin）。"""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = theme["fonts"]["latin"]
    run.font.color.rgb = _rgb(theme["colors"][color])
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        for e in rPr.findall(qn(tag)):
            rPr.remove(e)
        e = rPr.makeelement(qn(tag), {"typeface": theme["fonts"]["cjk"]})
        rPr.append(e)


def est_lines(text, font_pt, box_w_cm):
    """按字符宽度估算折行数：CJK≈1.0em，数字/字母≈0.52em，空格≈0.3em。"""
    w = 0.0
    for ch in text:
        w += font_pt * (1.0 if ord(ch) > 0x2E7F else (0.3 if ch == " " else 0.52))
    box_w_pt = box_w_cm / CM_PER_PT
    return max(1, math.ceil(w / max(box_w_pt, 1)))


class DeckBuilder:
    def __init__(self, theme):
        self.theme = theme
        self.w_cm, self.h_cm = SIZES[theme.get("size", "16:9")]
        self.prs = Presentation()
        self.prs.slide_width = Cm(self.w_cm)
        self.prs.slide_height = Cm(self.h_cm)
        self.blank = self.prs.slide_layouts[6]
        self.layout = []          # layout JSON sidecar 数据
        self._page_no = 0

    # ---------- 内部原语 ----------

    def _new_slide(self, layout_name):
        slide = self.prs.slides.add_slide(self.blank)
        self._page_no += 1
        rec = {"page": self._page_no, "layout": layout_name, "shapes": []}
        self.layout.append(rec)
        return slide, rec

    def _rect(self, slide, x, y, w, h, fill=None, line=False):
        sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
        sp.shadow.inherit = False
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = _rgb(self.theme["colors"][fill])
        if not line:
            sp.line.fill.background()
        return sp

    def _text(self, slide, rec, x, y, w, h, lines, size_key="body", bold=False,
              color="text", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              kind="body", line_spacing=1.22, space_after=6):
        """lines: list[str]；每个元素一个段落。"""
        tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        pt = self.theme["sizes"][size_key]
        for i, s in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            p.space_after = Pt(space_after)
            run = p.add_run()
            run.text = s
            _set_run_font(run, self.theme, pt, bold=bold, color=color)
        # fit 元数据（引擎侧估算，QA 侧校验）
        est = sum(est_lines(s, pt, w) for s in lines)
        need_h = est * pt * line_spacing * CM_PER_PT + 0.15
        rec["shapes"].append({
            "kind": kind, "box_cm": [round(x, 2), round(y, 2), round(w, 2), round(h, 2)],
            "font_pt": pt, "size_key": size_key, "paras": len(lines),
            "chars": sum(len(s) for s in lines),
            "est_lines": est, "est_need_h_cm": round(need_h, 2),
            "est_overflow": need_h > h + 0.05,
        })
        return tb

    def _footer(self, slide):
        s = self.theme["sizes"]["small"]
        tb = slide.shapes.add_textbox(Cm(self.w_cm - 3.0), Cm(self.h_cm - 1.0),
                                      Cm(2.4), Cm(0.7))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = str(self._page_no)
        _set_run_font(run, self.theme, s, color="muted")

    # ---------- 版式 ----------

    def add_cover(self, title, presenter="", major="", school="", advisor="",
                  date="", notes=""):
        slide, rec = self._new_slide("cover")
        self._rect(slide, 0, 0, self.w_cm, 2.2, fill="primary")
        self._rect(slide, 0, self.h_cm - 2.2, self.w_cm, 2.2, fill="primary")
        self._rect(slide, 2.0, 6.2, 3.2, 0.18, fill="accent")
        self._text(slide, rec, 2.0, 6.8, self.w_cm - 4.0, 4.2, [title],
                   size_key="cover_title", bold=True, kind="title")
        info = [s for s in (
            f"答辩人：{presenter}" if presenter else "",
            f"专业：{major}" if major else "",
            f"指导教师：{advisor}" if advisor else "",
            f"{school}　{date}".strip() if (school or date) else "",
        ) if s]
        self._text(slide, rec, 2.0, 11.6, self.w_cm - 4.0, 4.5, info,
                   size_key="cover_sub", color="muted", kind="meta",
                   line_spacing=1.4)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    def add_toc(self, items, title="目录", notes=""):
        slide, rec = self._new_slide("toc")
        self._rect(slide, 0, 0, 1.2, self.h_cm, fill="primary")
        self._text(slide, rec, 2.2, 1.4, 8.0, 2.0, [title], size_key="h1",
                   bold=True, color="primary", kind="title")
        lines = [f"{i + 1}　{t}" for i, t in enumerate(items)]
        box = self._rect(slide, 2.2, 3.8, self.w_cm - 4.4, self.h_cm - 5.4,
                         fill="light")
        box.shadow.inherit = False
        self._text(slide, rec, 3.0, 4.6, self.w_cm - 6.0, self.h_cm - 7.0,
                   lines, size_key="body", kind="body", line_spacing=1.5,
                   space_after=10)
        self._footer(slide)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    def add_section(self, no, title, notes=""):
        slide, rec = self._new_slide("section")
        self._rect(slide, 0, 0, self.w_cm, self.h_cm, fill="primary")
        self._text(slide, rec, 3.0, 5.4, 6.0, 4.0, [f"{no:02d}"],
                   size_key="section_no", bold=True, color="bg", kind="deco")
        self._text(slide, rec, 3.0, 9.6, self.w_cm - 6.0, 3.0, [title],
                   size_key="cover_title", bold=True, color="bg", kind="title")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    def add_content(self, title, bullets, notes="", kicker=""):
        """bullets: list[str] 或 list[(lvl, str)]，lvl∈{0,1}。"""
        slide, rec = self._new_slide("content")
        self._rect(slide, 0, 0, self.w_cm, 0.35, fill="primary")
        self._rect(slide, 2.0, 2.65, 1.6, 0.14, fill="accent")
        if kicker:
            self._text(slide, rec, 2.0, 0.9, self.w_cm - 4.0, 0.9, [kicker],
                       size_key="small", color="muted", kind="kicker")
        self._text(slide, rec, 2.0, 1.35 if not kicker else 1.75,
                   self.w_cm - 4.0, 1.6, [title], size_key="h1", bold=True,
                   kind="title")
        norm = [(0, b) if isinstance(b, str) else b for b in bullets]
        lines = [("　　" * lvl) + ("• " if lvl == 0 else "– ") + t
                 for lvl, t in norm]
        self._text(slide, rec, 2.2, 3.4, self.w_cm - 4.4, self.h_cm - 4.8,
                   lines, size_key="body", kind="body", space_after=8)
        self._footer(slide)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    def add_two_column(self, title, left_title, left_bullets, right_title,
                       right_bullets, notes=""):
        slide, rec = self._new_slide("two_column")
        self._rect(slide, 0, 0, self.w_cm, 0.35, fill="primary")
        self._text(slide, rec, 2.0, 1.35, self.w_cm - 4.0, 1.6, [title],
                   size_key="h1", bold=True, kind="title")
        col_w = (self.w_cm - 4.4 - 1.2) / 2
        for i, (ct, bl) in enumerate(((left_title, left_bullets),
                                      (right_title, right_bullets))):
            x = 2.2 + i * (col_w + 1.2)
            panel = self._rect(slide, x, 3.4, col_w, self.h_cm - 4.8,
                               fill="light")
            panel.shadow.inherit = False
            self._text(slide, rec, x + 0.5, 3.8, col_w - 1.0, 1.2, [ct],
                       size_key="body", bold=True, color="primary",
                       kind="col_title")
            self._text(slide, rec, x + 0.5, 5.1, col_w - 1.0,
                       self.h_cm - 6.6, ["• " + b for b in bl],
                       size_key="body", kind="body", space_after=8)
        self._footer(slide)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    def add_image_text(self, title, image_path, bullets, notes="",
                       image_side="left", caption=""):
        slide, rec = self._new_slide("image_text")
        self._rect(slide, 0, 0, self.w_cm, 0.35, fill="primary")
        self._text(slide, rec, 2.0, 1.35, self.w_cm - 4.0, 1.6, [title],
                   size_key="h1", bold=True, kind="title")
        img_w = (self.w_cm - 4.4) * 0.48
        img_h = self.h_cm - 6.2
        ix = 2.2 if image_side == "left" else self.w_cm - 2.2 - img_w
        tx = 2.2 if image_side == "right" else self.w_cm - 2.2 - img_w
        if not os.path.isfile(image_path):
            raise FileNotFoundError(f"图片不存在: {image_path}")
        slide.shapes.add_picture(image_path, Cm(ix), Cm(3.4), width=Cm(img_w))
        rec["shapes"].append({"kind": "image", "path": image_path,
                              "box_cm": [round(ix, 2), 3.4, round(img_w, 2),
                                         round(img_h, 2)]})
        if caption:
            self._text(slide, rec, ix, 3.4 + img_h + 0.15, img_w, 0.8,
                       [caption], size_key="small", color="muted",
                       align=PP_ALIGN.CENTER, kind="caption")
        self._text(slide, rec, tx, 3.4, self.w_cm - 4.4 - img_w,
                   self.h_cm - 4.8, ["• " + b for b in bullets],
                   size_key="body", kind="body", space_after=8)
        self._footer(slide)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    def add_closing(self, title="恳请各位老师批评指正", sub="谢谢聆听",
                    notes=""):
        slide, rec = self._new_slide("closing")
        self._rect(slide, 0, 0, self.w_cm, self.h_cm, fill="primary")
        self._text(slide, rec, 3.0, 7.2, self.w_cm - 6.0, 3.0, [title],
                   size_key="cover_title", bold=True, color="bg",
                   align=PP_ALIGN.CENTER, kind="title")
        self._text(slide, rec, 3.0, 11.0, self.w_cm - 6.0, 2.0, [sub],
                   size_key="cover_sub", color="bg", align=PP_ALIGN.CENTER,
                   kind="meta")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return self

    # ---------- 输出 ----------

    def save(self, out_pptx, layout_json=None):
        os.makedirs(os.path.dirname(os.path.abspath(out_pptx)), exist_ok=True)
        self.prs.save(out_pptx)
        if layout_json:
            with open(layout_json, "w", encoding="utf-8") as f:
                json.dump({"size": self.theme.get("size", "16:9"),
                           "pages": self.layout}, f,
                          ensure_ascii=False, indent=1)
        return out_pptx


if __name__ == "__main__":  # 冒烟：生成最小样片
    here = os.path.dirname(os.path.abspath(__file__))
    theme = load_theme(os.path.join(here, "..", "assets", "theme-default.yaml"))
    d = DeckBuilder(theme)
    d.add_cover(title="示例：基于 FMEA 的起落架收放系统故障分析",
                presenter="张三", major="飞行器维修工程技术",
                school="某职业技术大学", advisor="李老师", date="2026-06")
    d.add_toc(["研究背景与意义", "研究对象与方法", "FMEA 分析过程",
               "结果与讨论", "结论与展望"])
    d.add_section(1, "研究背景与意义")
    d.add_content("研究背景", [
        "起落架收放系统故障占机型机械类故障较高比例",
        "航线排故依赖经验，缺少系统化风险排序工具",
        (1, "FMEA 可将故障模式按 RPN 量化排序"),
    ], notes="开场 30 秒：从航线排故痛点引入。")
    d.add_two_column("方法对比", "FMEA", ["正向枚举功能失效模式", "输出 RPN 排序表"],
                     "FTA", ["逆向追溯顶事件原因", "适合单点深挖"])
    d.add_closing()
    out = os.path.join(here, "_smoke.pptx")
    d.save(out, os.path.join(here, "_smoke.layout.json"))
    print(f"OK: {out} ({len(d.layout)} pages)")
