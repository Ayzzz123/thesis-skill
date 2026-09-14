# -*- coding: utf-8 -*-
"""Test B：无学校 Word 模板 -> FORMAT_RECONSTRUCTION（不依赖模板母版，正常生成 DOCX）。
运行：python tests/test_b_format_reconstruction.py"""
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
import template_fidelity as TF
import docx_engine as E
from docx import Document
from docx.oxml.ns import qn

PASS, FAIL = 0, 0


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  PASS", name, extra)
    else:
        FAIL += 1
        print("  FAIL", name, extra)


def main():
    tmp = tempfile.mkdtemp(prefix="tf_testB_")
    empty_school = os.path.join(tmp, "school")
    os.makedirs(empty_school)
    out = os.path.join(tmp, "重建论文.docx")
    print("== Test B: FORMAT_RECONSTRUCTION ==")
    mode, tpath = TF.select_docx_mode(empty_school)
    check("模式选择=FORMAT_RECONSTRUCTION", mode == TF.MODE_FORMAT_RECONSTRUCTION)
    check("template_file=None", tpath is None)
    rec = TF.document_generation_record(mode, tpath)
    check("document_generation 记录", rec["mode"] == "format_reconstruction"
          and rec["template_file"] is None)
    # 重建路径：空白文档 + 通用原语（无需模板）
    d = Document()
    TF.add_h1(d, "第一章 绪论")
    TF.add_body(d, "这是重建模式下的正文段落，含引用[1]。")
    TF.add_h2(d, "1.1 测试节")
    TF.add_body(d, "二级正文内容。")
    TF.add_md_table(d, [["序号", "项"], ["1", "甲"]], cap_cn="表1-1 重建测试表",
                    cap_en="Tab.1-1 Reconstruction test table", font=10.5)
    TF.add_md_table(d, [["序号", "项"], ["2", "乙"]], cap_cn="表1-2 相邻表",
                    cap_en="Tab.1-2 Adjacent table", font=10.5)
    TF.add_h1(d, "参考文献")
    TF.add_body(d, "[1] 作者. 题名[J]. 期刊, 2024, 1(1): 1-2.")
    d.save(out)
    d2 = Document(out)
    caps = [p.text.strip() for p in d2.paragraphs if p.text.strip().startswith("表1-")]
    check("表题均存在", len(caps) == 2, str(caps))
    els = list(d2.element.body.iterchildren())
    tbl_idx = [i for i, e in enumerate(els) if e.tag == qn("w:tbl")]
    gaps_ok = all(any(els[j].tag == qn("w:p") for j in range(tbl_idx[k] + 1, tbl_idx[k + 1]))
                  for k in range(len(tbl_idx) - 1))
    check("相邻表间段落分隔（engine 自动保护）", gaps_ok)
    check("标题/正文/表内容齐全",
          any(p.text.strip().startswith("第一章") for p in d2.paragraphs)
          and any(p.text.strip().startswith("[1]") for p in d2.paragraphs))
    print(f"Test B 结果: PASS={PASS} FAIL={FAIL}")
    shutil.rmtree(tmp, ignore_errors=True)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
