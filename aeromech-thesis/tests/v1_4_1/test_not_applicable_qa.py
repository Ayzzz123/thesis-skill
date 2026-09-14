# -*- coding: utf-8 -*-
"""tests/v1_4_1/test_not_applicable_qa.py — 无模板 QA（NOT_APPLICABLE 模型）

验证：
1. 无模板（省略 --template / 传 none）运行时，模板对照项（TF-01~04、TF-19）输出 NOT_APPLICABLE（附 reason），
   绝不显示为 PASS；自含检查照常执行；退出码 0（无 FAIL，Delivery Gate 不因无模板而失败）。
2. 提供模板时：不出现 NOT_APPLICABLE（回到模板对照/规范驱动模式）。
3. --template 指向不存在文件 = 配置错误（退出码 3），不得静默降级为无模板。
运行：python tests/v1_4_1/test_not_applicable_qa.py
"""
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.abspath(os.path.join(HERE, "..", ".."))
SCRIPTS = os.path.join(SKILL, "scripts")
sys.path.insert(0, SCRIPTS)

TF = os.path.join(SCRIPTS, "tf_qa.py")
PASS, FAIL, SKIP = 0, 0, 0


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  PASS", name, extra)
    else:
        FAIL += 1
        print("  FAIL", name, extra)


def skip(name, why):
    global SKIP
    SKIP += 1
    print("  SKIPPED_WITH_REASON", name, "|", why)


def run_cli(args, cwd=None):
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    p = subprocess.run([sys.executable] + args, capture_output=True, text=True,
                       encoding="utf-8", env=env, cwd=cwd or SKILL)
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def build_fixture_docx(path):
    """构造一个自含检查可通过的最小 DOCX（无模板场景）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from PIL import Image
    d = Document()

    def set_run(r, size=12):
        rPr = r._element.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
        if rf.getparent() is None:
            rPr.append(rf)
        rf.set(qn("w:eastAsia"), "宋体")
        r.font.size = Pt(size)

    def body_p(text, line=400, first=200):
        p = d.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:line"), str(line)); sp.set(qn("w:lineRule"), "exact")
        pPr.append(sp)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLineChars"), str(first)); ind.set(qn("w:firstLine"), "0")
        pPr.append(ind)
        set_run(p.add_run(text))
        return p

    # 封面 1x1 表（占 tables[0]，物品封面自检用）
    t0 = d.add_table(rows=1, cols=1)
    cell = t0.cell(0, 0)
    cell.text = ""
    set_run(cell.paragraphs[0].add_run("本科生毕业论文"), 16)
    for lab in ["院（系）名称：", "专业名称：测试专业", "学生姓名：", "指导教师："]:
        p = cell.add_paragraph()
        set_run(p.add_run(lab), 14)
    # 前置
    d.add_paragraph("摘要", style="Heading 1")
    set_run(d.add_paragraph().add_run("本文测试摘要。"), 12)
    d.add_paragraph("关键词：测试；关键词")
    d.add_paragraph("ABSTRACT", style="Heading 1")
    set_run(d.add_paragraph().add_run("This is a test abstract."), 12)
    d.add_paragraph("KEY WORDS: test; keywords")
    # TOC 域
    p = d.add_paragraph()
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = r'TOC \o "1-3" \h \z \u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    # 正文
    d.add_paragraph("第1章 绪论", style="Heading 1")
    body_p("本文正文段落，用于字体字号与行距缩进检查。")
    # 图块（tables[1]：1x1 + drawing） + 图题
    png = os.path.join(os.path.dirname(path), "_fig.png")
    Image.new("RGB", (400, 260), "white").save(png)
    tb = d.add_table(rows=1, cols=1)
    tb.cell(0, 0).paragraphs[0].add_run().add_picture(png, width=Cm(8))
    cp = tb.cell(0, 0).add_paragraph()
    set_run(cp.add_run("图1-1　测试图"), 10.5)
    # 三线表（tables[2]）
    d.add_paragraph("表1-1　测试表")
    t2 = d.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "列A"; t2.cell(0, 1).text = "列B"
    t2.cell(1, 0).text = "值1"; t2.cell(1, 1).text = "值2"
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", "12"), ("bottom", "12")):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), sz)
        borders.append(el)
    t2._tbl.tblPr.append(borders)
    # 页眉
    hdr = d.sections[0].header
    hdr.paragraphs[0].text = "测试论文"
    # 后置
    d.add_paragraph("参考文献", style="Heading 1")
    for i in (1, 2, 3):
        body_p(f"[{i}] 测试文献条目。", line=240, first=0)
    d.add_paragraph("致谢", style="Heading 1")
    set_run(d.add_paragraph().add_run("感谢测试支持。"), 12)
    d.add_paragraph("附录A 测试附录", style="Heading 1")
    set_run(d.add_paragraph().add_run("附录内容。"), 12)
    d.save(path)


def build_fixture_template(path):
    """最小模板（非母版封面）：仅为验证『提供模板时不出现 N/A』。"""
    from docx import Document
    d = Document()
    d.add_paragraph("本科生毕业论文")
    d.add_paragraph("规范样例段落。")
    d.save(path)


def main():
    print("== test_not_applicable_qa ==")
    tmp = tempfile.mkdtemp(prefix="v141_na_")
    docx = os.path.join(tmp, "论文.docx")
    build_fixture_docx(docx)
    tpl = os.path.join(tmp, "template.docx")
    build_fixture_template(tpl)
    out1 = os.path.join(tmp, "qa_na")
    out2 = os.path.join(tmp, "qa_tpl")

    # ---- 1) 无模板运行：N/A + 无 FAIL + rc 0 ----
    rc, out = run_cli([TF, "--docx", docx, "--out", out1, "--refs-min", "3"])
    check("无模板运行 rc=0（无 FAIL）", rc == 0, f"rc={rc}")
    rep = open(os.path.join(out1, "tf-qa-report.md"), encoding="utf-8").read()
    for code in ("TF-01", "TF-02", "TF-03", "TF-04", "TF-19"):
        line = next((l for l in rep.split("\n") if l.startswith(f"- {code}")), "")
        check(f"{code} 记 NOT_APPLICABLE", "NOT_APPLICABLE" in line, line[:80])
        check(f"{code} 未伪造 PASS", "PASS" not in line, line[:60])
    check("报告含 reason（模板对照不可执行）", "模板对照不可执行" in rep)
    self_contained = [l for l in rep.split("\n") if l.startswith(("- TF-05", "- TF-06", "- TF-12"))]
    check("自含检查照常执行", all("NOT_APPLICABLE" not in l for l in self_contained))
    check("汇总行含 NOT_APPLICABLE 计数", "NOT_APPLICABLE=" in rep)

    # ---- 2) --template none 等价 ----
    rc2, _ = run_cli([TF, "--template", "none", "--docx", docx, "--out",
                      os.path.join(tmp, "qa_none"), "--refs-min", "3"])
    check("--template none 等价无模板（rc=0）", rc2 == 0, f"rc={rc2}")

    # ---- 3) 提供模板：检查行不出现 N/A ----
    rc3, _ = run_cli([TF, "--template", tpl, "--docx", docx, "--out", out2, "--refs-min", "3"])
    check("提供模板运行 rc∈{0,1}", rc3 in (0, 1), f"rc={rc3}")
    rep2 = open(os.path.join(out2, "tf-qa-report.md"), encoding="utf-8").read()
    na_lines = [l for l in rep2.split("\n") if l.startswith("- TF") and "NOT_APPLICABLE" in l]
    check("提供模板时不出现 NOT_APPLICABLE（检查行）", not na_lines, str(na_lines[:2]))

    # ---- 4) 模板路径不存在 = 配置错误（rc 3，不降级）----
    rc4, out4 = run_cli([TF, "--template", os.path.join(tmp, "missing.docx"), "--docx", docx,
                         "--out", os.path.join(tmp, "qa_err")])
    check("模板缺失 rc=3", rc4 == 3, f"rc={rc4}")
    check("模板缺失提示不静默降级", "不静默降级" in out4 or "配置错误" in out4)

    print(f"test_not_applicable_qa 结果: PASS={PASS} FAIL={FAIL} SKIP={SKIP}")
    shutil.rmtree(tmp, ignore_errors=True)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
