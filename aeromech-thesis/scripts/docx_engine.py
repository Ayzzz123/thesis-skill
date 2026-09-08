# -*- coding: utf-8 -*-
"""
docx_engine.py — DOCX 排版原语库（aeromech-thesis v1.0.0 delivery stabilization）

从已通过视觉验收的 build_docx 实现中抽取的与论文内容无关的排版函数。
供 Agent 在每次论文组装时按论文结构调用（生成项目专属 build_docx.py）。

已验证规则来源：thesis-test-1.0 39 页 PDF（页面流 PASS / 图+题注同页 PASS / 页码正常）。
"""
import os
import re
from PIL import Image

from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------- 常量 ----------------
PAGE_WIDTH_LANDSCAPE_CM = 25.7   # A4 横向 29.7 − 2×2.0 边距
BODY_WIDTH_PORTRAIT_CM = 15.5    # A4 纵向 21.0 − 3.0 − 2.5 边距

# PAGE_FLOW_OPTIMIZER：图块高度分级（推荐值，可按版式调整）
FIG_H_NORMAL = 9.5      # 常规图（框图/统计图/故障树）
FIG_H_TALL = 12.5       # 超高窄图 aspect<0.4（技术路线图等）
FIG_H_LANDSCAPE = 10.0  # 横向页内图

# 图片类型 → 目标宽度（正文可用宽度比例）
IMAGE_WIDTH_RULES = {
    "mermaid_flowchart": 0.80,
    "mermaid_system": 0.85,
    "matplotlib_stat": 0.90,
    "fta_tree": 0.88,
    "default": 0.85,
}


# ---------------- 字体/段落 ----------------
def set_font(run, cn="宋体", size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_para(doc, text="", size=12, cn="宋体", bold=False, align=None, indent=None,
             space_after=0, line=1.5, keep_with_next=False):
    """普通正文段落：默认不 keep_together，允许自然跨页。"""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(space_after)
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    if keep_with_next:
        pf.keep_with_next = True
    if text:
        parts = re.split(r"\*\*(.+?)\*\*", text)
        for i, seg in enumerate(parts):
            if not seg:
                continue
            run = p.add_run(seg)
            set_font(run, cn, size, bold or (i % 2 == 1))
    return p


HEAD_SIZES = {1: 16, 2: 14, 3: 12}


def add_heading(doc, text, level=1, page_break=False):
    """真实 Word Heading 标题；H1 可选 page_break_before（封面/摘要/目录/章节/方向切换）。"""
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(12 if level == 1 else 6)
    if page_break:
        pf.page_break_before = True
    pf.keep_with_next = True
    pf.keep_together = True
    run = p.add_run(text)
    set_font(run, "黑体", HEAD_SIZES[level], bold=False, color=RGBColor(0, 0, 0))
    return p


# ---------------- Section / 页码 ----------------
def setup_section(sec, landscape=False):
    """设置页面方向与边距。注意：不在本函数加页码（见 add_page_number_once）。"""
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
    else:
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.5)


def add_page_number_once(doc):
    """页码字段只添加一次到第一个 section；后续 section footer 继承（linked）。
    若每个 section 都加 PAGE 字段会产生 888/101010 重复页码。"""
    section = doc.sections[0]
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    set_font(run, "宋体", 9)
    # 后续新增 section 保持 footer 继承
    def link_footer(doc):
        for s in doc.sections[1:]:
            s.footer.is_linked_to_previous = True
    link_footer(doc)


# ---------------- 表格工具 ----------------
def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def row_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


def no_border_table(t):
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tblPr.append(borders)


def set_cell_margins_table(t, cm=0.08):
    tblPr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(cm * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def clean_cell(t):
    return t.replace("**", "").replace("`", "").strip()


def cell_borders(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, sz in edges.items():
        el = OxmlElement(f"w:{edge}")
        if sz:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def col_weights_for(header):
    """智能列宽：长文本列更宽，数值/短码列窄。"""
    weights = []
    for h in header:
        hl = clean_cell(h)
        if any(kw in hl for kw in ["S", "O", "D", "RPN", "编号", "FM", "分值"]):
            weights.append(0.55)
        elif len(hl) >= 8:
            weights.append(1.6)
        else:
            weights.append(1.0)
    return weights


def add_table(doc, header, rows, font_size=10.5, header_cn="黑体",
              repeat_header=True, row_keep=True, compact=False,
              landscape=False, col_weights=None, widths_cm=None):
    """三线表：tblHeader 表头重复 + cantSplit 行不拆分 + 可选紧凑（附录大表）。"""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_border_table(t)
    set_cell_margins_table(t, 0.04 if compact else 0.1)

    usable = PAGE_WIDTH_LANDSCAPE_CM if landscape else BODY_WIDTH_PORTRAIT_CM
    if widths_cm:
        widths = widths_cm
    else:
        if col_weights is None:
            col_weights = [1.0] * len(header)
        total_w = sum(col_weights)
        widths = [usable * w / total_w for w in col_weights]

    for ci, h in enumerate(header):
        c = t.cell(0, ci)
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(clean_cell(h))
        set_font(run, header_cn, font_size)
        cell_borders(c, {"top": 12, "bottom": 6, "left": None, "right": None})
        c.width = Cm(widths[ci])
    if repeat_header:
        row_repeat_header(t.rows[0])
    if row_keep:
        row_cant_split(t.rows[0])

    for ri, row in enumerate(rows):
        if row_keep:
            row_cant_split(t.rows[ri + 1])
        for ci in range(len(header)):
            c = t.cell(ri + 1, ci)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0 if compact else 1.15
            if compact:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
            run = p.add_run(clean_cell(row[ci] if ci < len(row) else ""))
            set_font(run, "宋体", font_size)
            cell_borders(c, {"left": None, "right": None, "top": None, "bottom": None})
            if ri == len(rows) - 1:
                cell_borders(c, {"bottom": 12})
            c.width = Cm(widths[ci])
    return t


def add_caption(doc, text, keep_with_next=False):
    """图题/表题段落。表题设 keep_with_next 与后续表格绑定。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, "宋体", 10.5)
    return p


# ---------------- FigureBlock（图+题注 不可拆分）----------------
def calc_image_width_cm(png, img_type="default", landscape=False):
    """PAGE_FLOW_OPTIMIZER 尺寸计算：按宽高比与分级高度上限求插入宽度。"""
    usable = PAGE_WIDTH_LANDSCAPE_CM if landscape else BODY_WIDTH_PORTRAIT_CM
    ratio = IMAGE_WIDTH_RULES.get(img_type, IMAGE_WIDTH_RULES["default"])
    w_target = usable * ratio
    try:
        with Image.open(png) as im:
            pw, ph = im.size
            aspect = pw / ph
    except Exception:
        aspect = 1.0
    if aspect < 0.4:
        fig_h_limit = FIG_H_TALL if not landscape else 12.0
    else:
        fig_h_limit = FIG_H_NORMAL if not landscape else FIG_H_LANDSCAPE
    w_by_h = fig_h_limit * aspect
    w = min(w_target, w_by_h)
    if aspect >= 0.4:
        if aspect >= 1.0:
            w = max(w, 7.5)
        else:
            w = max(w, 6.0)
    return min(w, usable)


def figure_block(doc, png_path, fig_caption, img_type="default",
                 anchor_el=None, landscape=False):
    """FigureBlock：无边框单列表格 1 行，行 cantSplit → 图片+图题整体不可跨页。
    返回表格 XML 元素（可用 anchor.addnext 插入到指定段落后）。"""
    w_cm = calc_image_width_cm(png_path, img_type, landscape)
    tb = doc.add_table(rows=1, cols=1)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_border_table(tb)
    set_cell_margins_table(tb, 0.05)
    row_cant_split(tb.rows[0])

    cell = tb.cell(0, 0)
    cell.width = Cm(w_cm)
    p_img = cell.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_together = True
    run = p_img.add_run()
    run.add_picture(png_path, width=Cm(w_cm))

    p_cap = cell.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.line_spacing = 1.3
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    cr = p_cap.add_run(fig_caption)
    set_font(cr, "宋体", 10.5)

    tbl_el = tb._tbl
    if anchor_el is not None:
        anchor_el.addnext(tbl_el)
    return tbl_el
