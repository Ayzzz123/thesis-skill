# -*- coding: utf-8 -*-
"""thesis_build.py — 统一 Thesis Build Pipeline（aeromech-thesis v1.6.0，Phase 4）

目标（指令 §五/§六）：项目不再为每篇论文手写 builder。输入=Build Contract + 研究上下文 +
学校上下文；输出=毕业论文.docx/.pdf + artifact-manifest.yaml。调度职责：
  LOAD PROJECT → LOAD SCHOOL CONTEXT → LOAD RESEARCH CONTEXT → LOAD FIGURE CONTEXT →
  LOAD CONTENT → BUILD DOCX → UPDATE TOC → REPAGINATE → EXPORT PDF → FINALIZE → QA → GATE
不重复实现 Template Fidelity / Figure QA / Research QA / PDF QA——组装一律走
template_fidelity + docx_engine 原语，链步骤一律调既有脚本。

Build Contract：`<root>/.aeromech/build-contract.yaml`
  project: {title, author, major, school}          # 必填 title
  content: {abstract_zh, abstract_zh_file, keywords, abstract_en, keywords_en,
            chapters: [rel paths], references_file, appendices: [{title, file}],
            ack_text | ack_file}
  research: {required: true|false}                  # 研究上下文（注册表）在场性由引擎判定，不伪造
  school_format: {template: rel|null, format_file}  # 缺 template→FORMAT_RECONSTRUCTION
  figures: [{figure_id, display, file, caption_cn, caption_en}]   # 或 figure_context 引用
  tables: [{display, caption_en}]（表体由章节 md 渲染，本域只提供英文题注）
  qa: {out: rel}  | output: {docx, pdf}
缺失语义（指令 §七）：必需内容缺 → ERROR（列出缺项）；可选域缺 → NOT_APPLICABLE（如实记录），
不得静默补假数据。

可重复性（指令 §十八）：content_identity(docx)=排除 docProps 与 _rels 后逐条目 sha256 归并；
pdf_text_identity(pdf)=逐页文本+内嵌图像 xref sha 归并。容器元数据（生成时间戳等）差异
不算内容漂移（区分 content identity vs container hash）。

用法：
  python thesis_build.py <root> validate [--json]
  python thesis_build.py <root> build            # 仅 DOCX 组装（两模式自动选择）
  python thesis_build.py <root> pipeline [--steps docx,toc,repaginate,pdf,finalize,qa]
  python thesis_build.py <root> manifest [--json]
退出码：0=成功；1=内容/前置失败（失败阶段+错误码+原因+建议动作已输出）；2=无 contract（NOT_APPLICABLE）；3=ERROR（环境/内部异常）
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import subprocess
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    import yaml
except ImportError:  # pragma: no cover
    print("需要 PyYAML（pip install pyyaml）")
    sys.exit(2)

import template_fidelity as TF
import docx_engine as E
import research_context as RC
import figure_iface as FI

SCRIPTS = os.path.dirname(os.path.abspath(__file__))
CONTRACT_PATH = ".aeromech/build-contract.yaml"
MANIFEST_PATH = ".aeromech/artifacts/build/artifact-manifest.yaml"
DEFAULT_DOCX = "毕业论文.docx"
DEFAULT_PDF = "毕业论文.pdf"


def _now():
    return datetime.datetime.now().astimezone().isoformat(timespec="seconds")


def contract_path(root):
    return os.path.join(root, *CONTRACT_PATH.split("/"))


def load_contract(root):
    p = contract_path(root)
    if not os.path.isfile(p):
        return None, None
    with open(p, encoding="utf-8") as f:
        raw = yaml.safe_load(f) or {}
    return raw, p


def _sha(path):
    if not path or not os.path.isfile(path):
        return None
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


_META_ENTRIES = re.compile(r"^(docProps/|word/_rels/|_rels/|customXml/)"
                           r"|docProps")


def content_identity(path):
    """docx 内容身份：zip 内非元数据条目的 (name, sha256) 归并哈希。
    排除 docProps（创建/修改时间、工具元数据）与关系文件——容器时间戳差异≠内容漂移。"""
    if not os.path.isfile(path):
        return None
    h = hashlib.sha256()
    with zipfile.ZipFile(path) as z:
        for name in sorted(z.namelist()):
            if name.startswith("docProps/") or name.endswith(".rels"):
                continue
            h.update(name.encode())
            h.update(hashlib.sha256(z.read(name)).hexdigest().encode())
    return h.hexdigest()


def pdf_text_identity(path):
    """PDF 内容身份：逐页文本 + 内嵌图像字节哈希（排除 CreationDate 等容器元数据）。"""
    if not os.path.isfile(path):
        return None
    try:
        import pymupdf as fitz
    except ImportError:
        return None
    h = hashlib.sha256()
    with fitz.open(path) as doc:
        for page in doc:
            h.update(page.get_text().encode("utf-8", "replace"))
            for x in page.get_images(full=True):
                try:
                    h.update(doc.extract_image(x[0])["image"])
                except Exception:
                    h.update(b"<img-fail>")
    return h.hexdigest()


# ---------------- 契约校验（缺失语义：必需缺=ERROR，可选缺=NOT_APPLICABLE） ----------------

def validate_contract(root, contract=None):
    c = contract if contract is not None else load_contract(root)[0]
    if c is None:
        return {"status": "NOT_APPLICABLE",
                "reason": "无 build-contract.yaml（旧项目经各自 builder 交付，v1.6 不接管文档域）",
                "errors": [], "na": ["document domain"], "warnings": []}
    errors, na, warns = [], [], []
    proj = c.get("project") or {}
    content = c.get("content") or {}
    if not proj.get("title"):
        errors.append("project.title 缺失（不得静默用占位题目）")
    ch = content.get("chapters") or []
    if not ch:
        errors.append("content.chapters 为空（没有章节就没有论文）")
    for f in ch:
        p = os.path.join(root, ".aeromech", f) if not os.path.isabs(f) else f
        if not os.path.isfile(p):
            errors.append(f"章节文件缺失：{f}")
        elif os.path.getsize(p) == 0:
            errors.append(f"章节文件为空：{f}")
    if content.get("abstract_zh") or content.get("abstract_zh_file"):
        if content.get("abstract_zh_file") and not os.path.isfile(
                os.path.join(root, ".aeromech", content["abstract_zh_file"])):
            errors.append("摘要文件缺失：" + content["abstract_zh_file"])
    else:
        na.append("abstract（未提供：跳过摘要页，交付报告披露）")
    if not content.get("references_file"):
        na.append("references（未提供）")
    elif not os.path.isfile(os.path.join(root, ".aeromech", content["references_file"])):
        errors.append("参考文献文件缺失：" + content["references_file"])
    for ap in content.get("appendices") or []:
        if not os.path.isfile(os.path.join(root, ".aeromech", ap.get("file", ""))):
            errors.append(f"附录文件缺失：{ap.get('file')}")
    # 研究域：注册表在场=OK；缺席=NOT_APPLICABLE（不报错、不伪造）
    ctx = RC.context(root)
    research_status = ctx["status"]
    if (c.get("research") or {}).get("required") and research_status != "OK":
        errors.append(f"research.required=true 但注册表状态={research_status}")
    # 学校域
    school_dir = os.path.join(root, "materials", "school")
    mode, tpl = TF.select_docx_mode(school_dir)
    sf = c.get("school_format") or {}
    if sf.get("template"):
        tp = os.path.join(root, sf["template"])
        if not os.path.isfile(tp):
            errors.append("school_format.template 指向不存在的文件：" + sf["template"])
        else:
            mode, tpl = TF.MODE_TEMPLATE_FIDELITY, tp
    elif not tpl:
        na.append("学校模板（未提供→FORMAT_RECONSTRUCTION，general_defaults，不冒充模板接管）")
    # 图域
    figs = c.get("figures") or []
    fig_na = False
    for fg in figs:
        p = os.path.join(root, fg.get("file", ""))
        if not fg.get("file") or not os.path.isfile(p):
            errors.append(f"figure 文件缺失：{fg.get('figure_id')} {fg.get('file')}")
    if not figs:
        fig_na = True
        na.append("figures（契约未列：正文含图占位但无文件时按 parse_md 诚实降级）")
    return {"status": "ERROR" if errors else "PASS", "errors": errors, "na": na,
            "warnings": warns, "docx_mode": mode, "template": tpl,
            "research_status": research_status, "figures_na": fig_na}


# ---------------- 组装（双模式，全走既有原语） ----------------

def _chapter_md(root, rel):
    p = os.path.join(root, ".aeromech", rel)
    with open(p, encoding="utf-8") as f:
        return f.read()


def _fig_maps(contract, root):
    """cap_map: 编号（parse_md 的 "1-1" 语义，display 去图/表前缀）→ (文件名, 中文题注)。
    英文题注按 display 原前缀分流：图→fig_en_map（FigureBlock 双语题注），表→en_map
    （表上方 Tab. 段）。编号在图/表各自序列里可同名（图3-1/表3-1），不得互相覆盖。
    表域（v1.6 test-8.0 接通）：contract.tables 条目只提供英文题注（表体由章节 md
    的表题行+管道行经 parse_md 渲染），无 file 要求。"""
    cap_map, en_map, fig_en_map = {}, {}, {}
    figdir = None
    for fg in contract.get("figures") or []:
        disp_raw = str(fg.get("display", ""))
        disp = re.sub(r"^[图表]\s*", "", disp_raw)
        fname = os.path.basename(fg.get("file", ""))
        if disp and fname:
            cap_map[disp] = (fname, fg.get("caption_cn", ""))
            figdir = os.path.dirname(os.path.join(root, fg["file"]))
        if disp and fg.get("caption_en"):
            (en_map if disp_raw.strip().startswith("表") else fig_en_map)[disp] = \
                fg["caption_en"]
    tb = contract.get("tables")
    entries = tb if isinstance(tb, list) else (tb or {}).get("captions") if isinstance(tb, dict) else None
    if isinstance(entries, list):
        for tg in entries:
            disp = re.sub(r"^[图表]\s*", "", str(tg.get("display", "")))
            if disp and tg.get("caption_en"):
                en_map[disp] = tg["caption_en"]
    return cap_map, en_map, figdir, fig_en_map


def build_docx(root, contract=None):
    """统一组装。返回 (out_docx, {mode, content_identity, sha256, sections})."""
    c = contract if contract is not None else load_contract(root)[0]
    v = validate_contract(root, c)
    if v["status"] == "ERROR":
        raise BuildError("契约校验失败", v["errors"])
    mode = v["docx_mode"]
    out_rel = (c.get("output") or {}).get("docx") or DEFAULT_DOCX
    out = out_rel if os.path.isabs(out_rel) else os.path.join(root, out_rel)
    content = c["content"]
    proj = c["project"]
    cap_map, en_map, figdir, fig_en_map = _fig_maps(c, root)

    if mode == TF.MODE_TEMPLATE_FIDELITY:
        doc = TF.open_master(v["template"], out)
        from docx.oxml.ns import qn
        sf = c.get("school_format") or {}
        cover_tables = sf.get("cover_tables")  # 保留的封面区表索引（如 [0,1]=封面+扉页）
        tbl_els = [el for el in doc.element.body.iterchildren() if el.tag == qn("w:tbl")]
        if cover_tables:
            idxs = [int(i) for i in cover_tables if 0 <= int(i) < len(tbl_els)]
            first_tbl = tbl_els[min(idxs)] if idxs else None
            keep_through = tbl_els[max(idxs)] if idxs else None
            # 前置裁剪：封面区之前是母版的规范说明页时（last_sectpr_anchor_before 找到
            # 封面前最后一个节锚点），把锚点段之前（含锚点段=前一节终结段）全部删除，
            # 文档以封面节为第一节。无锚点（封面即文档开头）→ 不动（v1.5 兼容）。
            if first_tbl is not None:
                front_anchor = TF.last_sectpr_anchor_before(doc, first_tbl)
                if front_anchor is not None:
                    # keep_anchor=False：锚点段的 sectPr 属于被剪掉的说明节，留着会产生
                    # 空白首节页；删除后文档以封面内容为第一节。
                    TF.trim_body_before(doc, front_anchor, keep_anchor=False)
        else:
            keep_through = tbl_els[0] if tbl_els else None  # v1.5 兼容：首表+首锚点
        anchor = TF.first_sectpr_anchor_after(doc, keep_through) if keep_through is not None else None
        if anchor is not None:
            TF.trim_body_after(doc, anchor)
            # 锚点段本身删除：其节属性与母版 body 级 sectPr 之间不留空节
            #（否则 _append_content 的 add_section 在其前再闭一节 → 空白页，
            #  test-8.0 P4 实测）。封面/扉页区的节由前部锚点界定，不受影响。
            TF.remove_anchor_paragraph(doc, anchor)
        TF.strip_all_pgnumtype(doc)  # 封面/扉页不显示页码；内容节页码由 _append_content 重建
        values = _cover_values(sf, proj, content)
        table_limit = (max(int(i) for i in cover_tables) + 1) if cover_tables else None
        filled = TF.fill_cover_fields(doc, values, table_limit=table_limit)
        texts = _fill_cover_texts(doc, sf.get("cover_text_fills"), proj, content)
        cover_info = {"filled": sorted(filled),
                      "unfilled": sorted(set(values) - set(filled)),
                      "text_filled": sorted(texts)}
        _append_content(root, doc, c, v, cap_map, en_map, figdir, roman_front=True,
                        fig_en_map=fig_en_map)
    else:
        from docx import Document
        doc = Document()
        sec = doc.sections[0]
        E.setup_section(sec)
        cover_info = None
        _append_content(root, doc, c, v, cap_map, en_map, figdir, roman_front=False,
                        fresh_cover=True, fig_en_map=fig_en_map)
        E.add_page_number_once(doc)
    doc.save(out)
    _mark_embedded(root, c, cap_map, figdir)
    return out, {"mode": mode, "template": v["template"], "sha256": _sha(out),
                 "content_identity": content_identity(out), "cover": cover_info,
                 "chapters": len(content.get("chapters") or [])}


def _cover_values(sf, proj, content):
    """school_format.cover_fields: {封面标签: 取值来源}；来源支持 project.<key>/
    content.<key>/字面量。无配置→v1.5 兼容硬编码（题/专业）。"""
    fields = sf.get("cover_fields")
    if not fields:
        return {"题    目": proj["title"], "专    业": proj.get("major", "")}
    out = {}
    for label, src in fields.items():
        s = str(src)
        if s.startswith("project."):
            v = proj.get(s[8:])
        elif s.startswith("content."):
            v = content.get(s[8:])
        else:
            v = s
        if v is not None and str(v).strip():
            out[label] = str(v)
    return out


def _fill_cover_texts(doc, fills, proj=None, content=None):
    """封面正文段/表内段的字面占位替换（如日期"二○XX年X月"——封面与扉页各一处，
    同一值全部填充）。值支持 project.<k>/content.<k>/字面量（与 cover_fields 同一
    取值语法）。跨 run 时折叠重写（保留首 run 格式）。未命中的如实进 unfilled 披露。"""
    filled = {}
    if not fills:
        return filled
    proj = proj or {}
    content = content or {}
    paras = list(doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for cell in r.cells:
                paras.extend(cell.paragraphs)
    for pattern, src in fills.items():
        s = str(src)
        if s.startswith("project."):
            value = proj.get(s[8:])
        elif s.startswith("content."):
            value = content.get(s[8:])
        else:
            value = s
        if value is None or not str(value).strip():
            continue
        n_hit = 0
        for p in paras:
            if pattern in p.text:
                done = False
                for run in p.runs:
                    if pattern in run.text:
                        run.text = run.text.replace(pattern, str(value))
                        done = True
                        break
                if not done and p.runs:
                    full = "".join(r.text or "" for r in p.runs)
                    if pattern in full:
                        p.runs[0].text = full.replace(pattern, str(value))
                        for r in p.runs[1:]:
                            r.text = ""
                        done = True
                if done:
                    n_hit += 1
        if n_hit:
            filled[pattern] = str(value)
    return filled


def _mark_embedded(root, c, cap_map, figdir):
    """诚实记录：仅当章节 md 的占位行会被 parse_md 同一正则命中、且图文件在场时，
    才记 EMBEDDED。只为已有 VALIDATED/GENERATED 状态的图补记录，不凭空创建。"""
    if not cap_map or not figdir:
        return
    fid_by_name = {os.path.basename(str(f.get("file", ""))): str(f.get("figure_id"))
                   for f in (c.get("figures") or []) if f.get("file")}
    hit_keys = set()
    md_re1 = re.compile(r"^（图(\d+-\d+)[^）]*（[^）]*）[^）]*）$")
    md_re2 = re.compile(r"^（图(\d+-\d+)[^）]*）$")
    for f in c.get("content", {}).get("chapters") or []:
        try:
            with open(os.path.join(root, ".aeromech", f), encoding="utf-8") as fh:
                for ln in fh:
                    m = md_re1.match(ln.strip()) or md_re2.match(ln.strip())
                    if m and m.group(1) in cap_map:
                        hit_keys.add(m.group(1))
        except OSError:
            continue
    for disp in hit_keys:
        val = cap_map.get(disp)
        fname = val[0] if isinstance(val, tuple) else None
        if not fname or not os.path.isfile(os.path.join(figdir, fname)):
            continue
        fid = fid_by_name.get(fname)
        if not fid:
            continue
        cur, _ = FI.current_status(root, fid)
        if cur in ("VALIDATED", "GENERATED"):
            FI.record(root, fid, "EMBEDDED",
                      artifact=os.path.join(figdir, fname).replace("\\", "/"),
                      reason="统一构建 figure_block 插入（章节占位行命中）")


class BuildError(Exception):
    def __init__(self, msg, details=None):
        super().__init__(msg + ("; " + "; ".join(details) if details else ""))
        self.details = details or []


def _content_text(root, content, key):
    """内联字段优先；否则读 *_file（validate 已查存在性）。不伪造缺省文本。"""
    f = content.get(key + "_file")
    if f:
        p = os.path.join(root, ".aeromech", f)
        with open(p, encoding="utf-8") as fh:
            return fh.read().strip()
    return content.get(key)


def _append_content(root, doc, c, v, cap_map, en_map, figdir, roman_front,
                    fresh_cover=False, fig_en_map=None):
    content = c["content"]
    proj = c["project"]
    if fresh_cover:
        p = doc.add_paragraph()
        TF.ppr(p, line=600, rule="exact", jc="center", before=3600)
        TF.add_para_runs(p, proj["title"], ea="黑体", sz=22)
        for line in [proj.get("school"), proj.get("major"), proj.get("author")]:
            q = doc.add_paragraph()
            TF.ppr(q, line=400, rule="exact", jc="center", before=240)
            TF.add_para_runs(q, line or "", ea="宋体", sz=14)
        doc.add_page_break()
    # 摘要（各自独立成页，AbstractBlock 规则）
    zh = _content_text(root, content, "abstract_zh")
    zh_file = content.get("abstract_zh_file")
    if zh or zh_file:
        if roman_front:
            sec = TF.add_section_continue(doc, fmt="upperRoman")
            TF.set_pgnum(doc.sections[-1], fmt="upperRoman", start=1)
            TF.footer_page_field(sec)
        TF.add_h1(doc, "摘  要", page_break=False)
        TF.add_body(doc, zh or "")
        kw = doc.add_paragraph()
        TF.ppr(kw, line=400, rule="exact")
        TF.add_para_runs(kw, "关键词：" + "；".join(content.get("keywords") or []),
                         ea="宋体", sz=12)
        en = _content_text(root, content, "abstract_en")
        if en:
            TF.add_h1(doc, "ABSTRACT", page_break=True)
            TF.add_body(doc, en)
            kw = doc.add_paragraph()
            TF.ppr(kw, line=400, rule="exact")
            TF.add_para_runs(kw, "KEY WORDS: " + "; ".join(content.get("keywords_en") or []),
                             ea="宋体", sz=12, bold=True)
    # 目录域（原生 TOC；页码由 update_toc COM 刷新）
    sec = TF.add_section_continue(doc, fmt="upperRoman") if roman_front else doc.add_section()
    if roman_front:
        TF.set_pgnum(doc.sections[-1], fmt=None)
        TF.footer_page_field(doc.sections[-1])
    tp = doc.add_paragraph()
    TF.ppr(tp, line=400, rule="exact", jc="center")
    TF.add_para_runs(tp, "目  录", ea="黑体", sz=16)
    TF.toc_field_para(doc)
    # 正文
    sec = TF.add_section_continue(doc, fmt="decimal") if roman_front else doc.add_section()
    if roman_front:
        TF.set_pgnum(doc.sections[-1], fmt="decimal", start=1)
    # 页眉（学校规范：从正文开始；契约 school_format.header，
    # 支持 {title}=论文题目插值；缺省不加页眉=按规范由 QA 判）
    hdr = ((c.get("school_format") or {}).get("header") or "").strip()
    if hdr:
        TF.header_text(doc.sections[-1], hdr.replace("{title}", proj.get("title", "")))
    for f in content.get("chapters") or []:
        TF.parse_md(doc, _chapter_md(root, f), fig_dir=figdir,
                    cap_map=cap_map or None, en_map=en_map or None,
                    fig_en_map=fig_en_map or None)
    # 参考文献（md 列表→段；无则跳过+披露）
    rf = content.get("references_file")
    if rf and os.path.isfile(os.path.join(root, ".aeromech", rf)):
        TF.add_h1(doc, "参考文献", page_break=True)
        with open(os.path.join(root, ".aeromech", rf), encoding="utf-8") as fh:
            for ln in fh.read().splitlines():
                s = ln.strip().lstrip("-").strip()
                if s:
                    TF.add_body(doc, s)
    # 附录 / 致谢
    for ap in content.get("appendices") or []:
        TF.add_h1(doc, ap.get("title", "附录"), page_break=True)
        TF.parse_md(doc, _chapter_md(root, ap["file"]))
    ack = content.get("ack_text") or ""
    if content.get("ack_file"):
        ack = _chapter_md(root, content["ack_file"])
    if ack:
        TF.add_h1(doc, "致  谢", page_break=True)
        TF.add_body(doc, ack)


# ---------------- manifest（指令 §八） ----------------

def write_manifest(root, entries, build_id=None, meta=None):
    build_id = build_id or f"B-{datetime.datetime.now():%Y%m%d%H%M%S}"
    items = []
    for e in entries:
        path = e["path"]
        ap = path if os.path.isabs(path) else os.path.join(root, path)
        items.append({"artifact": e["artifact"], "path": path.replace("\\", "/"),
                      "sha256": _sha(ap), "content_identity": e.get("content_identity"),
                      "producer": e.get("producer", "thesis_build"),
                      "timestamp": e.get("timestamp") or _now(),
                      "status": e.get("status", "OK"), "reason": e.get("reason", "")})
    data = {"build_id": build_id, "ts": _now(), "meta": meta or {}, "artifacts": items}
    p = os.path.join(root, *MANIFEST_PATH.split("/"))
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        yaml.safe_dump(data, f, allow_unicode=True, sort_keys=False)
    return build_id, items, p


def load_manifest(root):
    p = os.path.join(root, *MANIFEST_PATH.split("/"))
    if not os.path.isfile(p):
        return None
    with open(p, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def check_manifest(root):
    """Resume 对接（指令 §八）：路径存在性 + sha256 一致性。"""
    man = load_manifest(root)
    if man is None:
        return {"status": "NOT_APPLICABLE", "reason": "无 artifact-manifest（旧项目/未走统一构建）",
                "missing": [], "changed": []}
    missing, changed = [], []
    for it in man.get("artifacts") or []:
        ap = os.path.join(root, it["path"])
        if not os.path.isfile(ap):
            missing.append(it["path"])
        elif _sha(ap) != it.get("sha256"):
            changed.append(it["path"])
    return {"status": "ERROR" if missing or changed else "PASS",
            "missing": missing, "changed": changed, "build_id": man.get("build_id")}


# ---------------- 流水线（指令 §六/§十四） ----------------

def _step(root, name, cmd, timeout=600, allow_na=False):
    """运行既有脚本子进程。返回步记录 {step, rc, status, output_tail}。"""
    rec = {"step": name, "ts": _now()}
    try:
        p = subprocess.run(cmd, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=timeout,
                           cwd=root)
        rc = p.returncode
    except FileNotFoundError as e:
        rc, out = 3, f"脚本缺失: {e}"
    except subprocess.TimeoutExpired:
        rc, out = 3, f"超时 {timeout}s"
    tail = ((p.stdout or "") + "\n" + (p.stderr or "")).strip().splitlines()[-3:] if rc != 3 else [out]
    if rc == 0:
        st = "PASS"
    elif allow_na and rc == 2:
        st = "NOT_APPLICABLE"
    elif rc == 3:
        st = "ERROR"
    else:
        st = "FAIL"
    rec.update({"rc": rc, "status": st, "output_tail": " | ".join(tail)})
    return rec


ALL_STEPS = ["docx", "toc", "repaginate", "pdf", "finalize", "qa"]


def pipeline(root, steps=None):
    c, cp = load_contract(root)
    if c is None:
        return {"status": "NOT_APPLICABLE", "reason": f"无 {CONTRACT_PATH}（旧项目各自 builder，v1.6 不接管）",
                "steps": []}
    root = root
    steps = steps or ALL_STEPS
    rec = {"contract": cp, "steps": [], "artifacts": [], "status": None}
    out_docx = (c.get("output") or {}).get("docx") or DEFAULT_DOCX
    out_pdf = (c.get("output") or {}).get("pdf") or DEFAULT_PDF
    abs_docx = out_docx if os.path.isabs(out_docx) else os.path.join(root, out_docx)
    abs_pdf = out_pdf if os.path.isabs(out_pdf) else os.path.join(root, out_pdf)
    qa_out = os.path.join(root, ".aeromech", (c.get("qa") or {}).get("out") or "artifacts/qa")
    os.makedirs(qa_out, exist_ok=True)
    tpl = (c.get("school_format") or {}).get("template")
    tpl_abs = os.path.join(root, tpl) if tpl else TF.select_docx_mode(
        os.path.join(root, "materials", "school"))[1]
    # 模板 PDF（母版导出的视觉基准）：与 tpl 同路径的 .pdf，或契约显式 template_pdf；
    # 供 cover_align/cover_fill/color/page 的 --template-pdf 使用（缺则这些步按依赖跳过）。
    tpl_pdf = None
    if tpl_abs:
        cand = [(c.get("school_format") or {}).get("template_pdf"),
                os.path.splitext(tpl_abs)[0] + ".pdf"]
        for x in cand:
            if x:
                xp = os.path.join(root, x) if not os.path.isabs(x) else x
                if os.path.isfile(xp):
                    tpl_pdf = xp
                    break
    py = sys.executable

    def log_step(s):
        rec["steps"].append(s)
        print(f"[{s['status']:<15}] {s['step']}  rc={s.get('rc')}"
              + (f"  {s['output_tail'][:120]}" if s.get("output_tail") else ""))

    if "docx" in steps:
        try:
            out, info = build_docx(root, c)
            s = {"step": "docx", "status": "PASS", "rc": 0, "ts": _now(),
                 "output_tail": f"mode={info['mode']} chapters={info['chapters']}"}
            rec["artifacts"].append({"artifact": "docx", "path": out_docx,
                                     "content_identity": info["content_identity"],
                                     "producer": "thesis_build"})
        except BuildError as e:
            s = {"step": "docx", "status": "FAIL", "rc": 1, "ts": _now(),
                 "output_tail": str(e)}
            rec["failed_stage"] = "docx"
            log_step(s)
            for rest in [x for x in steps if x != "docx"]:
                log_step({"step": rest, "status": "SKIPPED_WITH_REASON",
                          "reason": "docx 构建失败，下游不执行（不假装成功）",
                          "ts": _now()})
            rec["status"] = "FAIL"
            rec["verdict"] = {"failed_stage": "docx", "error_code": "BUILD_CONTRACT",
                              "reason": str(e),
                              "suggested_action": "补齐契约缺失项后重跑 build"}
            return rec
        log_step(s)
    depend_ok = all(x["status"] in ("PASS", "NOT_APPLICABLE") for x in rec["steps"])
    if "toc" in steps:
        if depend_ok:
            s = _step(root, "toc", [py, os.path.join(SCRIPTS, "update_toc.py"), root])
            if s["status"] == "ERROR":
                s["output_tail"] = "Word COM 不可用/环境失败（非内容结论）"
            log_step(s)
        else:
            s = {"step": "toc", "status": "SKIPPED_WITH_REASON", "reason": "docx 步骤失败"}
            log_step(s)
    if "repaginate" in steps and depend_ok:
        log_step(_step(root, "repaginate",
                       [py, os.path.join(SCRIPTS, "repaginate_tables.py"), root]))
    if "pdf" in steps:
        if depend_ok:
            s = _step(root, "pdf", [py, os.path.join(SCRIPTS, "export_pdf.py"), root],
                      timeout=900)
            if s["status"] == "PASS" and os.path.isfile(abs_pdf):
                rec["artifacts"].append({"artifact": "pdf", "path": out_pdf,
                                         "content_identity": pdf_text_identity(abs_pdf),
                                         "producer": "export_pdf"})
            log_step(s)
        else:
            log_step({"step": "pdf", "status": "SKIPPED_WITH_REASON",
                      "reason": "上游失败", "rc": None, "ts": _now()})
    if "finalize" in steps and depend_ok:
        log_step(_step(root, "finalize",
                       [py, os.path.join(SCRIPTS, "finalize_metadata.py"), root]))
    if "qa" in steps:
        docx_ok = os.path.isfile(abs_docx)
        pdf_ok = os.path.isfile(abs_pdf)
        chain = [
            ("pdf_qa", [py, os.path.join(SCRIPTS, "pdf_qa.py"), root], pdf_ok),
            ("visual_regression", [py, os.path.join(SCRIPTS, "visual_regression.py"), root], pdf_ok),
            ("tf_qa", [py, os.path.join(SCRIPTS, "tf_qa.py"),
                       "--template", tpl_abs or "none", "--docx", abs_docx,
                       "--pdf", abs_pdf if pdf_ok else "none", "--out", qa_out], docx_ok),
            ("cover_fidelity", [py, os.path.join(SCRIPTS, "cover_fidelity.py"),
                                "--template", tpl_abs or "none", "--docx", abs_docx,
                                "--pdf", abs_pdf if pdf_ok else "none", "--out", qa_out],
             docx_ok and bool(tpl_abs)),
            ("content_purity", [py, os.path.join(SCRIPTS, "content_purity_qa.py"),
                                "--docx", abs_docx, "--pdf", abs_pdf if pdf_ok else "none",
                                "--out", qa_out], docx_ok),
            # v1.6 test-8.0 接通：delivery gate 的 FORMAT_REPORTS 列 10 项格式 QA，
            # 此前 chain 只跑 5 项（cover_align/cover_fill/color/page/table 从未执行
            # →报告永远缺席）。全部为既有脚本，按各自 CLI 挂进链。
            ("cover_align", [py, os.path.join(SCRIPTS, "cover_align_qa.py"),
                             "--pdf", abs_pdf,
                             *(["--template-pdf", tpl_pdf] if tpl_pdf else []),
                             "--out", qa_out], bool(tpl_pdf) and pdf_ok),
            ("cover_fill", [py, os.path.join(SCRIPTS, "cover_fill_qa.py"),
                            "--pdf", abs_pdf,
                            *(["--template-pdf", tpl_pdf] if tpl_pdf else []),
                            "--out", qa_out], bool(tpl_pdf) and pdf_ok),
            ("color_fidelity", [py, os.path.join(SCRIPTS, "color_fidelity_qa.py"),
                                "--docx", abs_docx, "--pdf", abs_pdf,
                                *(["--template-docx", tpl_abs, "--template-pdf", tpl_pdf]
                                  if tpl_pdf and tpl_abs else []),
                                "--out", qa_out], docx_ok and pdf_ok),
            ("page_fidelity", [py, os.path.join(SCRIPTS, "page_fidelity_qa.py"),
                               "--docx", abs_docx, "--pdf", abs_pdf,
                               *(["--template-pdf", tpl_pdf] if tpl_pdf else []),
                               *(["--en-title", c["content"]["abstract_en"]]
                                 if c.get("content", {}).get("abstract_en") else []),
                               "--out", qa_out], docx_ok and pdf_ok),
            ("table_readability", [py, os.path.join(SCRIPTS, "table_readability_qa.py"),
                                   "--docx", abs_docx, "--pdf", abs_pdf,
                                   "--out", qa_out], docx_ok and pdf_ok),
        ]
        figdir = os.path.join(root, ".aeromech", "artifacts", "figures")
        if os.path.isdir(figdir):
            tmin = str((c.get("qa") or {}).get("tables_min", 15))
            chain.append(("figure_table", [py, os.path.join(SCRIPTS, "figure_table_qa.py"),
                                           "--docx", abs_docx, "--pdf", abs_pdf,
                                           "--figdir", figdir, "--tables-min", tmin,
                                           "--out", qa_out], pdf_ok))
            chain.append(("graph_quality", [py, os.path.join(SCRIPTS, "graph_quality_qa.py"),
                                            "--docx", abs_docx, "--pdf", abs_pdf,
                                            "--figdir", figdir, "--out", qa_out], pdf_ok))
        for name, cmd, needed in chain:
            if not needed:
                log_step({"step": name, "status": "SKIPPED_WITH_REASON",
                          "reason": "依赖产物不存在（无模板/无 PDF/无 DOCX）", "ts": _now()})
                continue
            log_step(_step(root, name, cmd, timeout=600, allow_na=name == "tf_qa"))
    # 图生命周期域证据
    lst = FI.lifecycle_state(root)
    if lst["has_plan"]:
        bad = {f: x["status"] for f, x in lst["figures"].items()
               if x["status"] in ("REJECTED", "NEEDS_HUMAN_REVIEW")}
        rec["figures"] = {"status": "FAIL" if any(v == "REJECTED" for v in bad.values())
                          else ("NEEDS_HUMAN_REVIEW" if bad else "PASS"),
                          "bad": bad}
    else:
        rec["figures"] = {"status": "NOT_APPLICABLE"}
    hard_fail = [s for s in rec["steps"] if s["status"] in ("FAIL", "ERROR")]
    rec["status"] = "ERROR" if any(s["status"] == "ERROR" for s in hard_fail) \
        else ("FAIL" if hard_fail else "PASS")
    rec["manifest_build_id"], _, rec["manifest_path"] = write_manifest(
        root, rec["artifacts"],
        meta={"steps": {s["step"]: s["status"] for s in rec["steps"]},
              "status": rec["status"]})
    if rec["status"] != "PASS":
        fs = hard_fail[0]
        rec["verdict"] = {"failed_stage": fs["step"], "error_code": f"STEP_{fs['status']}",
                          "rc": fs.get("rc"), "reason": fs.get("output_tail"),
                          "suggested_action": ("修复环境后重跑该步" if fs["status"] == "ERROR"
                                               else "按报告修复内容后重跑 pipeline")}
    return rec


# ---------------- CLI ----------------

def main(argv=None):
    ap = argparse.ArgumentParser(description="Unified Thesis Build（v1.6）")
    ap.add_argument("root")
    sub = ap.add_subparsers(dest="cmd", required=True)
    v = sub.add_parser("validate"); v.add_argument("--json", action="store_true")
    b = sub.add_parser("build"); b.add_argument("--json", action="store_true")
    pl = sub.add_parser("pipeline")
    pl.add_argument("--steps", default=",".join(ALL_STEPS))
    pl.add_argument("--json", action="store_true")
    m = sub.add_parser("manifest"); m.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    root = os.path.abspath(a.root)
    try:
        if a.cmd == "validate":
            res = validate_contract(root)
            print(json.dumps(res, ensure_ascii=False, indent=1) if a.json else
                  f"validate: {res['status']}" + "".join(f"\n  - {e}" for e in res["errors"])
                  + "".join(f"\n  · N/A: {x}" for x in res["na"]))
            return {"PASS": 0, "NOT_APPLICABLE": 2, "ERROR": 1}[res["status"]]
        if a.cmd == "build":
            c, _ = load_contract(root)
            if c is None:
                print("无契约：NOT_APPLICABLE（旧项目各自 builder）")
                return 2
            root = root
            out, info = build_docx(root, c)
            write_manifest(root, [{"artifact": "docx", "path": os.path.relpath(out, root),
                                   "content_identity": info["content_identity"]}],
                           meta={"mode": info["mode"]})
            print(json.dumps({"docx": out, **info}, ensure_ascii=False, default=str))
            return 0
        if a.cmd == "pipeline":
            steps = [s for s in a.steps.split(",") if s]
            res = pipeline(root, steps=steps)
            print(json.dumps(res, ensure_ascii=False, indent=1, default=str) if a.json else "")
            if res["status"] == "NOT_APPLICABLE":
                return 2
            if res["status"] == "PASS":
                return 0
            vdict = res.get("verdict", {})
            print(f"PIPELINE {res['status']}: 失败阶段={vdict.get('failed_stage')} "
                  f"错误码={vdict.get('error_code')} 原因={vdict.get('reason')} "
                  f"建议={vdict.get('suggested_action')}")
            return 3 if res["status"] == "ERROR" else 1
        if a.cmd == "manifest":
            chk = check_manifest(root)
            man = load_manifest(root) or {}
            out = {**chk, "build_id": man.get("build_id"),
                   "artifacts": man.get("artifacts") or []}
            print(json.dumps(out, ensure_ascii=False, indent=1) if a.json else
                  f"manifest: {chk['status']} build_id={chk.get('build_id') or man.get('build_id')}"
                  f" missing={chk.get('missing')} changed={chk.get('changed')}")
            return 0 if chk["status"] in ("PASS", "NOT_APPLICABLE") else 1
    except BuildError as e:
        print(f"BUILD FAIL: {e}")
        return 1
    except Exception as e:
        print(f"ERROR: {type(e).__name__}: {e}")
        return 3
    return 0


if __name__ == "__main__":
    sys.exit(main())
